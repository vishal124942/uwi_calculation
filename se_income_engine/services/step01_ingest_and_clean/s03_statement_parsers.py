"""M1 Parser - statement extraction into canonical AccountInput objects (Section 3, 24.1 M1).

Supported sources:
* Bank-statement analyser XLSX export (``Transactions`` sheet with Date/Narration/Credit/Debit/Balance).
* Generic CSV (date, narration, credit, debit, balance [, reference]).

Bank/vendor specific columns are persisted in ``parser_fields`` only (Section 3.3).
"""

from __future__ import annotations

import csv
import hashlib
import re
from datetime import date, datetime
from pathlib import Path

from se_income_engine.common.money import to_money
from se_income_engine.contracts.account import AccountInput, Transaction
from se_income_engine.contracts.enums import AccountType, Direction, LegalOwnerType, SourceType
from se_income_engine.versions import COMPONENT_VERSIONS

_HEADER_RE = re.compile(r"Bank\s*:\s*(?P<bank>[A-Z0-9 &]+?)\s+A/c No\.?:\s*(?P<acct>\S+)\s*-\s*(?P<name>.+)$", re.I)
_CORE_COLUMNS = {"Date", "Narration", "Credit", "Debit", "Balance"}


def stable_transaction_id(account_id: str, *parts) -> str:
    payload = "|".join([account_id, *[str(p) for p in parts]])
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:24]


def _to_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _build_transaction(account_id: str, source_line: int, page: int | None, txn_date, narration, credit, debit,
                       balance, reference, extra: dict[str, str]) -> Transaction | None:
    booking_date = _to_date(txn_date)
    credit_amt = to_money(credit)
    debit_amt = to_money(debit)
    has_credit = credit_amt is not None and credit_amt != 0
    has_debit = debit_amt is not None and debit_amt != 0
    if booking_date is None or has_credit == has_debit:
        return None
    direction = Direction.CREDIT if has_credit else Direction.DEBIT
    amount = abs(credit_amt if has_credit else debit_amt)
    raw = None if narration is None else str(narration)
    bal = to_money(balance)
    return Transaction(
        transaction_id=stable_transaction_id(account_id, source_line, booking_date, direction.value, amount, raw, bal),
        account_id=account_id,
        booking_date=booking_date,
        direction=direction,
        amount=amount,
        balance_after=bal,
        raw_narration=raw,
        source_reference=None if reference in (None, "") else str(reference),
        parser_fields=extra,
        source_page=page,
        source_line=source_line,
        parse_confidence=1.0,
    )


def parse_analyser_xlsx(
    path: str | Path,
    customer_id: str,
    account_id: str | None = None,
    account_type: AccountType = AccountType.SAVINGS,
    legal_owner_type: LegalOwnerType = LegalOwnerType.UNKNOWN,
    currency: str = "INR",
) -> AccountInput:
    import openpyxl

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Transactions" not in wb.sheetnames:
        raise ValueError(f"{path.name}: no 'Transactions' sheet found")
    rows = list(wb["Transactions"].iter_rows(values_only=True))

    metadata: dict[str, str] = {"source_file": path.name}
    if rows and rows[0] and rows[0][0]:
        m = _HEADER_RE.search(str(rows[0][0]))
        if m:
            metadata.update(bank=m.group("bank").strip(), masked_account=m.group("acct").strip(),
                            holder_name=m.group("name").strip())
    for r in rows[:5]:
        if r and r[0] == "DOC ID" and len(r) > 1:
            metadata["doc_id"] = str(r[1])

    header_idx = next(
        (i for i, r in enumerate(rows) if r and r[0] == "Date" and "Narration" in r),
        None,
    )
    if header_idx is None:
        raise ValueError(f"{path.name}: transaction header row not found")
    header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[header_idx])]
    col = {h: i for i, h in enumerate(header)}

    acct_id = account_id or metadata.get("doc_id") or path.stem
    txns: list[Transaction] = []
    source_records = 0
    rejected = 0
    for offset, r in enumerate(rows[header_idx + 1:]):
        if r is None or all(v in (None, "") for v in r):
            continue
        source_records += 1
        extra = {
            h: ("" if r[i] is None else str(r[i]))
            for h, i in col.items()
            if h not in _CORE_COLUMNS and i < len(r)
        }
        txn = _build_transaction(
            acct_id,
            source_line=header_idx + 2 + offset,
            page=None,
            txn_date=r[col["Date"]],
            narration=r[col["Narration"]],
            credit=r[col["Credit"]],
            debit=r[col["Debit"]],
            balance=r[col["Balance"]] if "Balance" in col else None,
            reference=r[col["Cheque/Ref No"]] if "Cheque/Ref No" in col else None,
            extra=extra,
        )
        if txn is None:
            rejected += 1
            continue
        txns.append(txn)

    metadata["source_record_count"] = str(source_records)
    metadata["rejected_record_count"] = str(rejected)
    dates = [t.booking_date for t in txns]
    return AccountInput(
        account_id=acct_id,
        customer_id=customer_id,
        account_type=account_type,
        legal_owner_type=legal_owner_type,
        currency=currency,
        statement_start=min(dates) if dates else date.today(),
        statement_end=max(dates) if dates else date.today(),
        source_type=SourceType.CSV,
        parser_version=COMPONENT_VERSIONS["parser"],
        transactions=txns,
        bank_code=metadata.get("bank"),
        parser_metadata=metadata,
    )


def parse_generic_csv(
    path: str | Path,
    customer_id: str,
    account_id: str,
    account_type: AccountType = AccountType.SAVINGS,
    legal_owner_type: LegalOwnerType = LegalOwnerType.UNKNOWN,
    currency: str = "INR",
) -> AccountInput:
    path = Path(path)
    txns: list[Transaction] = []
    source_records = 0
    rejected = 0
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        lower = {h.lower().strip(): h for h in reader.fieldnames or []}

        def pick(*names):
            return next((lower[n] for n in names if n in lower), None)

        c_date, c_narr = pick("date", "txn date", "booking date"), pick("narration", "description", "particulars")
        c_cr, c_dr = pick("credit", "deposit", "cr"), pick("debit", "withdrawal", "dr")
        c_bal, c_ref = pick("balance", "closing balance"), pick("reference", "ref no", "cheque/ref no")
        if c_date is None or c_cr is None or c_dr is None:
            raise ValueError(f"{path.name}: requires date, credit and debit columns")
        used = {c_date, c_narr, c_cr, c_dr, c_bal, c_ref}
        for line_no, row in enumerate(reader, start=2):
            source_records += 1
            extra = {k: (v or "") for k, v in row.items() if k not in used}
            txn = _build_transaction(account_id, line_no, None, row.get(c_date), row.get(c_narr) if c_narr else None,
                                     row.get(c_cr), row.get(c_dr), row.get(c_bal) if c_bal else None,
                                     row.get(c_ref) if c_ref else None, extra)
            if txn is None:
                rejected += 1
                continue
            txns.append(txn)
    dates = [t.booking_date for t in txns]
    return AccountInput(
        account_id=account_id,
        customer_id=customer_id,
        account_type=account_type,
        legal_owner_type=legal_owner_type,
        currency=currency,
        statement_start=min(dates) if dates else date.today(),
        statement_end=max(dates) if dates else date.today(),
        source_type=SourceType.CSV,
        parser_version="generic_csv_parser_v1.0.0",
        transactions=txns,
        parser_metadata={"source_file": path.name, "source_record_count": str(source_records),
                         "rejected_record_count": str(rejected)},
    )
